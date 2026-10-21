"""
文档解析：上传的文件 → 可入库的文本段落
=========================================
支持 4 种格式：pdf / docx / md / txt

docx 取的是**正文顺序上的全部块**（段落 + 表格），表格每行渲染成 `单元格 | 单元格`。
不用 `document.paragraphs` 的原因：它只返回正文直接段落，表格在另一棵树上，
用它等于把"以表格为主的文档"解析成空文档 —— 状态还是 ready，检索却永远查不到。

为什么中间要有「段落」（ParsedSegment）这一层，而不是直接返回一整段字符串：
  切片时需要知道「这段文字来自原文件的哪里」，检索命中后才能告诉用户出处。
  PDF 有天然的页边界 → 每页一个段落，带 page_ref="p.3"
  其他格式没有页概念 → 整篇一个段落，page_ref=None
  这一段小设计就是后面「引用定位到页码」的地基。

为什么解析要丢进线程（asyncio.to_thread）：
  pypdf / python-docx 都是**同步阻塞**实现，一份 PDF 解析几百毫秒到几秒。
  直接 await 调用等于在事件循环上跑同步代码 —— 会把整个进程卡住。
  处理方式和 D9 的 embedding 同源：CPU/IO 密集的同步代码 → 丢线程池。
"""

import asyncio
import io
from dataclasses import dataclass
from pathlib import PurePosixPath

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

# ---- 上传限制（在路由层做便宜校验，不解析内容）----
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB


@dataclass
class ParsedSegment:
    """解析出来的一段文本（PDF 里的一段 = 一页）。"""

    text: str
    page_ref: str | None = None   # "p.3"；无页概念时为 None


def check_upload(filename: str, data: bytes) -> str:
    """
    上传前置校验（只看扩展名和大小，不解析内容），返回 file_type。

    为什么要和「解析」分开：
    这类错误是**用户输入问题**，应该在请求阶段直接 400 回给用户，
    不必先建一条文档记录再让它变 failed。
    而真正解析时才发现的错误（坏 PDF）走另一条路 —— 落成 failed 状态。
    """
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = " / ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"不支持的文件类型「{extension or '无扩展名'}」，仅支持 {supported}")
    if not data:
        raise ValueError("文件内容为空")
    if len(data) > MAX_FILE_SIZE:
        raise ValueError(f"文件过大（{len(data) / 1024 / 1024:.1f}MB），上限 {MAX_FILE_SIZE // 1024 // 1024}MB")
    return extension.lstrip(".")


async def parse_file(filename: str, data: bytes) -> list[ParsedSegment]:
    """
    按扩展名选择解析器，把文件内容转成文本段落。

    解析本身是同步阻塞的，所以整体丢进线程池执行（asyncio.to_thread）。
    """
    extension = PurePosixPath(filename).suffix.lower()
    parser = _PARSERS.get(extension)
    if parser is None:                                  # 正常不会走到（路由层已校验）
        raise ValueError(f"没有对应的解析器：{extension}")

    segments = await asyncio.to_thread(parser, data)
    if not segments:
        # 关键：不要静默返回空列表。
        # 空列表会被下游当成「切片数为 0」，文档看着是 ready 却检索不到任何内容 ——
        # 这种「沉默的失败」比直接报错难查得多。
        raise ValueError("未从文件中提取到任何文本（可能是扫描件/纯图片文件）")
    return segments


# ------------------------------------------------------------
# 各格式解析器（同步实现，由 parse_file 丢进线程池调用）
# ------------------------------------------------------------


def _parse_pdf(data: bytes) -> list[ParsedSegment]:
    """PDF 按页解析，每页一个段落并记下页码。"""
    reader = PdfReader(io.BytesIO(data))
    segments: list[ParsedSegment] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:                                        # 空白页跳过，不产生空切片
            segments.append(ParsedSegment(text=text, page_ref=f"p.{page_no}"))
    return segments


def _iter_docx_blocks(container, parent):
    """
    按**文档顺序**逐块产出 Paragraph / Table。

    为什么要自己走 XML，而不是 `document.paragraphs`：
    后者只返回**正文直接段落**。表格里的段落在
    `table.rows[].cells[].paragraphs` 这棵完全不同的树上，
    于是"以表格为主的文档"（报销标准、权限矩阵、价格表）会被解析成
    一段没有表格内容的文本 —— 文档照样 ready、切片数照样正常，
    但**检索永远查不到表格里的字**，而且全程没有任何报错。

    三种子节点要认：
      w:p    普通段落
      w:tbl  表格
      w:sdt  内容控件（一层壳，正文包在 w:sdtContent 里）。
             企业模板的"填写区域"常用它，不钻进去又是一处静默丢内容。
    """
    for child in container.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_docx_blocks(content, parent)


def _table_lines(table: Table) -> list[str]:
    """
    表格 → 每行一条 `单元格 | 单元格` 的文本。

    三个必须处理的细节：
    ① **合并单元格**：`row.cells` 会把同一个格子重复返回 N 次（横向合并了几列就几次），
       不去重就会入库成「住宿 | 住宿 | 每晚 600 元」这种重复噪声。
       按底层 `tc` 元素的身份去重（python-docx 对合并格子返回的是同一个 tc）。
    ② **空行**：整行都没文字就跳过，否则切片里会堆出一排 ` | `。
    ③ **单元格里还有段落 / 嵌套表格**：段落用空格连，嵌套表格递归展开成行、用分号连。
       单元格不是"一个字符串"，它自己也是一棵块级树。
    """
    lines: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:                       # 合并单元格：同一个 tc 被重复返回
                continue
            seen.add(key)
            cells.append(_cell_text(cell))
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _cell_text(cell) -> str:
    """单元格内的文字：段落用空格连，嵌套表格递归成行再用分号连。"""
    parts: list[str] = []
    for block in _iter_docx_blocks(cell._tc, cell):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                parts.append(text)
        else:
            nested = _table_lines(block)
            if nested:
                parts.append("；".join(nested))
    return " ".join(parts)


def _parse_docx(data: bytes) -> list[ParsedSegment]:
    """
    DOCX 按文档顺序取正文段落 + 表格，拼成一段文本（docx 没有页的概念）。

    顺序很重要：「一、标准如下表」→ 表格 → 「以上自 2026 年 1 月起执行」
    拼错顺序，模型会把结论和它对应的表格对不上号。
    """
    document = DocxDocument(io.BytesIO(data))
    lines: list[str] = []

    for block in _iter_docx_blocks(document.element.body, document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(text)
        else:
            lines.extend(_table_lines(block))

    text = "\n".join(lines)          # 过滤空行，避免切片里全是换行
    return [ParsedSegment(text=text)] if text else []


def _parse_plain_text(data: bytes) -> list[ParsedSegment]:
    """md / txt：解码后即用。"""
    text = _decode(data).strip()
    return [ParsedSegment(text=text)] if text else []


def _decode(data: bytes) -> str:
    """
    解码纯文本。

    为什么要试两种编码：中文用户手里的 txt 大量来自 Windows 记事本，
    默认是 GBK。只试 utf-8 会在这些文件上直接抛 UnicodeDecodeError。
    """
    for encoding in ("utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法解码文本内容（已尝试 utf-8 / gbk）")


_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".md": _parse_plain_text,
    ".txt": _parse_plain_text,
}
