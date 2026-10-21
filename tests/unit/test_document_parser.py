"""
单元：上传校验与文档解析（D12）
================================
两个入口，职责刻意分开：
- `check_upload` 只看**扩展名和大小**（不解析内容）→ 用户输入错误应当场 400
- `parse_file`   真去解析 → 坏文件走"落成 failed 状态"那条路

这个分工本身就是考点：便宜校验放前面，贵的解析放后面。
"""

import io

import pytest
from docx import Document as DocxDocument

from app.services.document_parser import (
    MAX_FILE_SIZE,
    SUPPORTED_EXTENSIONS,
    check_upload,
    parse_file,
)

# ============================================================
# 一、前置校验（同步、纯逻辑）
# ============================================================


@pytest.mark.parametrize("filename", ["a.pdf", "a.docx", "a.md", "a.txt", "A.PDF", "报告.MD"])
def test_supported_extensions_pass(filename):
    """扩展名大小写不敏感（Windows 导出的大写扩展名很常见）。"""
    file_type = check_upload(filename, b"content")
    assert file_type in {"pdf", "docx", "md", "txt"}


@pytest.mark.parametrize("filename", ["a.zip", "a.exe", "noext", "a.md.zip"])
def test_unsupported_extensions_rejected(filename):
    with pytest.raises(ValueError, match="不支持的文件类型"):
        check_upload(filename, b"content")


def test_empty_content_rejected():
    """空文件当场拒 —— 否则它会变成一条"解析出 0 切片"的记录，看着成功了其实没内容。"""
    with pytest.raises(ValueError, match="内容为空"):
        check_upload("a.txt", b"")


def test_oversize_rejected():
    with pytest.raises(ValueError, match="过大"):
        check_upload("a.txt", b"x" * (MAX_FILE_SIZE + 1))


def test_size_limit_boundary_is_inclusive():
    """恰好等于上限要**放行**（边界属于合法侧），否则用户会被一个说不清理由的 400 拦住。"""
    assert check_upload("a.txt", b"x" * MAX_FILE_SIZE) == "txt"


def test_supported_set_is_what_the_docs_claim():
    """PRD 说支持 4 种格式，这里把清单钉住（防有人加了格式忘改文档）。"""
    assert SUPPORTED_EXTENSIONS == {".pdf", ".docx", ".md", ".txt"}


# ============================================================
# 二、真解析（async —— 解析被丢进线程池，见模块文档）
# ============================================================


async def test_markdown_is_parsed_as_single_segment():
    segments = await parse_file("note.md", "第一行\n第二行".encode())
    assert len(segments) == 1
    assert segments[0].text == "第一行\n第二行"
    assert segments[0].page_ref is None  # md 没有"页"这个概念


async def test_gbk_encoded_text_is_decoded():
    """
    ★ 中文用户的 txt 大量来自 Windows 记事本，默认编码是 **GBK**。
    只试 utf-8 会在这些文件上直接抛 UnicodeDecodeError —— 而用户完全不知道为什么。
    """
    segments = await parse_file("gbk.txt", "中文内容测试".encode("gbk"))
    assert segments[0].text == "中文内容测试"


async def test_docx_paragraphs_are_joined_and_blank_lines_dropped():
    buffer = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph("第一段")
    document.add_paragraph("")
    document.add_paragraph("第二段")
    document.save(buffer)

    segments = await parse_file("a.docx", buffer.getvalue())
    assert segments[0].text == "第一段\n第二段"
    assert segments[0].page_ref is None


async def test_whitespace_only_file_raises_instead_of_returning_empty():
    """
    ★ 关键设计：**宁可报错，不要静默返回空**。
    空列表会被下游当成"切片数为 0"，文档看着是 ready 却检索不到任何东西 ——
    "沉默的失败"比直接报错难查得多。
    """
    with pytest.raises(ValueError, match="未从文件中提取到任何文本"):
        await parse_file("blank.md", "   \n\n\t  ".encode())


async def test_unknown_extension_reaching_parser_raises():
    """路由层已拦过一次，但解析器自己也要兜底（防有人绕过路由直接调）。"""
    with pytest.raises(ValueError, match="没有对应的解析器"):
        await parse_file("a.zip", b"PK\x03\x04")


# ============================================================
# 三、DOCX 表格（块 3：修静默丢数据）
# ============================================================


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_docx_table_content_is_kept_in_document_order():
    """
    ★ 这条用例对应一个真实的静默故障：
    `document.paragraphs` 只返回正文直接段落，表格在另一棵树上，
    于是"以表格为主"的文档入库后等于空文档 —— ready、切片数正常、检索永远查不到。

    同时锁住**顺序**：段落 → 表格 → 段落，错序会让"结论"和它对应的表对不上。
    """
    document = DocxDocument()
    document.add_paragraph("一、报销标准如下表所示。")
    table = document.add_table(rows=3, cols=2)
    rows = [("项目", "上限"), ("市内交通", "每天 50 元"), ("住宿", "每晚上限 600 元")]
    for row_index, (key, value) in enumerate(rows):
        table.cell(row_index, 0).text = key
        table.cell(row_index, 1).text = value
    document.add_paragraph("以上标准自 2026 年 1 月起执行。")

    segments = await parse_file("a.docx", _docx_bytes(document))
    text = segments[0].text

    assert "市内交通 | 每天 50 元" in text
    assert "住宿 | 每晚上限 600 元" in text
    assert text.index("一、报销标准") < text.index("市内交通") < text.index("以上标准")


async def test_docx_table_only_document_is_not_an_empty_parse():
    """整篇只有一个表格：旧实现会得到空文本 → 抛「未提取到任何文本」，用户以为文件坏了。"""
    document = DocxDocument()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "E104"

    segments = await parse_file("only-table.docx", _docx_bytes(document))
    assert segments[0].text == "型号 | E104"


async def test_docx_merged_cells_are_not_repeated():
    """合并单元格时 `row.cells` 会把同一个格子返回多次，不去重就会入库成重复文本。"""
    document = DocxDocument()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "住宿标准"

    segments = await parse_file("merged.docx", _docx_bytes(document))
    assert segments[0].text.count("住宿标准") == 1


async def test_docx_empty_table_does_not_produce_separator_noise():
    """整行都是空单元格时跳过该行 —— 否则切片里会堆出一排 ` | `。"""
    document = DocxDocument()
    document.add_paragraph("正文")
    document.add_table(rows=2, cols=2)

    segments = await parse_file("empty-table.docx", _docx_bytes(document))
    assert segments[0].text == "正文"


async def test_docx_nested_table_is_flattened_into_its_cell():
    """单元格里还能嵌表格，单元格本身也是一棵块级树。"""
    document = DocxDocument()
    outer = document.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "一级"
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "子项"
    inner.cell(0, 1).text = "值"

    segments = await parse_file("nested.docx", _docx_bytes(document))
    assert "一级 | 子项 | 值" in segments[0].text
